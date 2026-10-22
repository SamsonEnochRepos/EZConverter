"""Conformance: every advertised capability is executed against a real input with the real engine.

Outputs are checked by the product's own verifier (inside the worker) AND by independent tools wherever
one exists. Results are appended to tests/.conformance-results/*.jsonl; scripts/build_format_matrix.py turns
them into the verified-capability manifest and the format matrix documentation.
"""

from __future__ import annotations

import gzip
import io
import json
import os
import re
import subprocess
import time
import zipfile
from functools import cache
from pathlib import Path

import pytest
from fixtures.generate import ARCHIVE_FILES, SENTINEL, fixture

from ezconverter.registry import default_registry
from ezconverter.sandbox import Limits
from ezconverter.workers.launcher import run_worker

pytestmark = pytest.mark.conformance
REG = default_registry()
RESULTS = Path(__file__).resolve().parent.parent / ".conformance-results"
TEXT_SOURCES = {
    "docx",
    "doc",
    "odt",
    "rtf",
    "txt",
    "html",
    "md",
    "xlsx",
    "xls",
    "ods",
    "csv",
    "pptx",
    "ppt",
    "odp",
    "pdf",
}
IM_NAME = {
    "png": "PNG",
    "jpeg": "JPEG",
    "webp": "WEBP",
    "avif": "AVIF",
    "gif": "GIF",
    "tiff": "TIFF",
    "bmp": "BMP",
    "ico": "ICO",
}
AUDIO_CODEC = {
    "mp3": {"mp3"},
    "wav": {"pcm_s16le"},
    "flac": {"flac"},
    "aac": {"aac"},
    "m4a": {"aac"},
    "ogg": {"vorbis"},
    "opus": {"opus"},
    "aiff": {"pcm_s16be"},
}
VIDEO_CODEC = {
    "mp4": {"h264", "mpeg4", "hevc", "av1", "mpeg2video"},
    "mkv": {"h264", "vp8", "vp9", "mpeg4", "mpeg2video", "hevc", "av1"},
    "mov": {"h264", "mpeg4", "hevc"},
    "avi": {"mpeg4", "h264", "mpeg2video"},
    "webm": {"vp8", "vp9", "av1"},
    "mpeg": {"mpeg2video"},
}
AUDIO_FMTS = set(AUDIO_CODEC)
VIDEO_FMTS = set(VIDEO_CODEC)


def _cases():
    out = []
    for e in sorted(REG.all(), key=lambda e: e.capability.id):
        marks = []
        if not e.availability.ok:
            marks.append(pytest.mark.skip(reason=f"engine unavailable: {e.availability.reason}"))
        out.append(pytest.param(e, id=e.capability.id, marks=marks))
    return out


@cache
def _inspect(path: str) -> dict:
    wd = Path(path).parent / ".inspect" / Path(path).name
    r = run_worker(
        {"op": "inspect", "path": path, "filename": Path(path).name},
        wd,
        Limits(wall_seconds=60),
    )
    return r["detection"]


def _source_for(cap, tmp_path: Path) -> Path:
    if cap.source == "gz":
        payload = fixture(cap.target).read_bytes()
        p = tmp_path / "payload.gz"
        p.write_bytes(gzip.compress(payload, mtime=0))
        return p
    return fixture(cap.source)


def run_cap(entry, src: Path, tmp_path: Path, options: dict | None = None) -> dict:
    cap, adapter = entry.capability, entry.adapter
    det = (
        _inspect(str(src))
        if not src.parent == tmp_path
        else run_worker(
            {"op": "inspect", "path": str(src), "filename": src.name},
            tmp_path / "insp",
            Limits(wall_seconds=60),
        )["detection"]
    )
    assert det["format"] == cap.source, det
    opts = adapter.validate_options(cap, options)
    from ezconverter.engines.base import InputInfo

    adapter.assess(
        cap,
        InputInfo(src, src.name, cap.source, src.stat().st_size, det["details"]),
        opts,
    )  # must not raise for valid input
    wd = tmp_path / "job"
    req = {
        "op": "convert",
        "input": {
            "path": str(src),
            "filename": src.name,
            "format": cap.source,
            "size": src.stat().st_size,
            "details": det["details"],
        },
        "steps": [{"capability_id": cap.id, "options": opts}],
        "output_dir": str(wd / "out"),
        "work_dir": str(wd / "work"),
        "output_stem": "result",
        "limits": {"wall_seconds": 280, "cpu_seconds": 600},
    }
    return run_worker(req, wd, Limits(wall_seconds=300, cpu_seconds=600))


# ---------------------------------------------------------------------------------- independent checks


def sh(*argv: str) -> str:
    return subprocess.run(argv, capture_output=True, text=True, timeout=120, check=True).stdout


def ff(path: Path) -> dict:
    return json.loads(
        sh(
            "ffprobe",
            "-v",
            "error",
            "-show_format",
            "-show_streams",
            "-of",
            "json",
            str(path),
        )
    )


def independent(cap, out: Path, src: Path, tmp_path: Path) -> None:  # noqa: C901
    s, t = cap.source, cap.target
    if t in IM_NAME:
        name, w, h = sh("identify", "-format", "%m %w %h", f"{out}[0]").split()
        # ImageMagick reports BMP variants as BMP2/BMP3 and PNG-compressed ICO frames as PNG; the ICO header is checked directly.
        assert (
            (name.startswith(IM_NAME[t]) or (t == "ico" and name in ("ICO", "PNG", "BMP")))
            and int(w) > 0
            and int(h) > 0
        )
        if t == "ico":
            assert out.read_bytes()[:4] == b"\x00\x00\x01\x00"
        if s in IM_NAME or s == "svg":
            vals = sh(
                "convert",
                f"{out}[0]",
                "-background",
                "white",
                "-alpha",
                "remove",
                "-format",
                "%[fx:int(255*p{0.1*w,0.5*h}.r)],%[fx:int(255*p{0.1*w,0.5*h}.b)],%[fx:int(255*p{0.85*w,0.5*h}.r)],%[fx:int(255*p{0.85*w,0.5*h}.b)]",
                "info:",
            ).split(",")
            lr, lb, rr, rb = map(int, vals)
            assert lr > 150 and lb < 100 and rb > 150 and rr < 100, (
                f"left/right colour mismatch {vals}"
            )  # red | blue
        else:
            assert float(sh("convert", f"{out}[0]", "-format", "%[fx:mean]", "info:")) < 0.9995, (
                "rendered page is blank"
            )
    elif t in AUDIO_FMTS:
        info = ff(out)
        a = [x for x in info["streams"] if x["codec_type"] == "audio"]
        assert a and a[0]["codec_name"] in AUDIO_CODEC[t]
        assert not [
            x
            for x in info["streams"]
            if x["codec_type"] == "video" and not x.get("disposition", {}).get("attached_pic")
        ]
        assert abs(float(info["format"]["duration"]) - 2.0) < 0.35
        vol = subprocess.run(
            ["ffmpeg", "-i", str(out), "-af", "volumedetect", "-f", "null", "-"],
            capture_output=True,
            text=True,
        ).stderr
        mean = float(re.search(r"mean_volume: (-?[\d.]+) dB", vol).group(1))
        assert mean > -45, f"output is (near) silent: {mean} dB"
    elif t in VIDEO_FMTS:
        info = ff(out)
        v = [x for x in info["streams"] if x["codec_type"] == "video"]
        assert v and v[0]["codec_name"] in VIDEO_CODEC[t] and v[0]["width"] == 320 and v[0]["height"] == 240
        assert any(x["codec_type"] == "audio" for x in info["streams"])
        assert abs(float(info["format"]["duration"]) - 2.0) < 0.5
    elif t == "pdf":
        pages = int(re.search(r"Pages:\s+(\d+)", sh("pdfinfo", str(out))).group(1))
        assert pages >= 1
        if s in TEXT_SOURCES:
            # -raw keeps line-end hyphens (the default mode de-hyphenates, so a wrap right after "-" would hide it);
            # whitespace is dropped because the layout engine may wrap the sentinel across lines.
            assert "".join(SENTINEL.split()) in "".join(sh("pdftotext", "-raw", str(out), "-").split())
    elif t in ("docx", "odt", "rtf", "html", "md"):
        fmt = {"md": "gfm"}.get(t, t)
        assert SENTINEL in sh("pandoc", "-f", fmt, "-t", "plain", str(out)) if s in TEXT_SOURCES else True
    elif t == "txt":
        text = out.read_text("utf-8")
        assert SENTINEL in text if s in TEXT_SOURCES else True
    elif t in ("doc", "xls", "ppt"):
        raw = out.read_bytes()
        assert SENTINEL.encode() in raw or SENTINEL.encode("utf-16-le") in raw, (
            "sentinel missing from legacy binary output"
        )
    elif t == "xlsx":
        import openpyxl

        ws = openpyxl.load_workbook(out).active
        assert any(SENTINEL in str(c.value) for row in ws.iter_rows() for c in row)
    elif t == "pptx":
        from pptx import Presentation

        assert SENTINEL in " ".join(
            sh_.text_frame.text for sl in Presentation(out).slides for sh_ in sl.shapes if sh_.has_text_frame
        )
    elif t in ("ods", "odp"):
        with zipfile.ZipFile(out) as z:
            assert SENTINEL in z.read("content.xml").decode()
    elif t == "csv" and s in TEXT_SOURCES | {"json"}:
        import csv

        rows = list(csv.reader(io.StringIO(out.read_text("utf-8"), newline="")))
        assert any(SENTINEL in c for r in rows for c in r) and len({len(r) for r in rows}) == 1
    elif t == "zip":
        assert "OK" in sh("unzip", "-tq", str(out)) or "No errors" in sh("unzip", "-t", str(out))
        for n, b in ARCHIVE_FILES.items():
            assert subprocess.run(["unzip", "-p", str(out), n], capture_output=True).stdout == b
    elif t in ("tar", "tgz"):
        for n, b in ARCHIVE_FILES.items():
            assert subprocess.run(["tar", "-xOf", str(out), n], capture_output=True).stdout == b
    elif t == "7z":
        import py7zr

        dest = tmp_path / "x7"
        with py7zr.SevenZipFile(out) as z:
            z.extractall(dest)
        for n, b in ARCHIVE_FILES.items():
            assert (dest / n).read_bytes() == b
    elif t == "gz":
        assert gzip.decompress(out.read_bytes()) == src.read_bytes()  # byte-preserving
    elif s == "gz":
        assert out.read_bytes() == gzip.decompress(src.read_bytes())
    elif t == "json" and s in ("yaml", "csv", "xml"):
        data = json.loads(out.read_text())
        if s == "yaml":
            import yaml

            assert data == yaml.safe_load(src.read_text())
        elif s == "csv":
            import csv

            rows = list(csv.DictReader(io.StringIO(src.read_text(), newline="")))
            assert data == rows
        else:
            assert data["catalog"]["item"][0] == {
                "@id": "1",
                "name": "widget",
                "note": SENTINEL,
            }
            assert data["catalog"]["item"][1] == {"@id": "2", "name": "gadget"}
    elif t == "yaml":
        import yaml

        assert yaml.safe_load(out.read_text()) == json.loads(src.read_text())
    elif t == "xml":
        from xml.etree import ElementTree as ET

        root = ET.fromstring(out.read_text())
        assert root.get("ez-format") == "typed-json-1" and root.get("type") == "array" and len(root) == 2
        assert root[0].find("note").text == SENTINEL and root[0].find("tags").get("type") == "null"


def record(entry, status: str, seconds: float, extra: dict | None = None) -> None:
    RESULTS.mkdir(exist_ok=True)
    cap = entry.capability
    line = {
        "id": cap.id,
        "status": status,
        "seconds": round(seconds, 2),
        "engine": cap.engine_id,
        "version": entry.availability.version,
        **(extra or {}),
    }
    with open(RESULTS / f"{os.getpid()}.jsonl", "a") as fh:
        fh.write(json.dumps(line) + "\n")


@pytest.mark.parametrize("entry", _cases())
def test_capability_converts_real_input_and_output_is_correct(entry, tmp_path):
    cap = entry.capability
    t0 = time.time()
    try:
        src = _source_for(cap, tmp_path)
        res = run_cap(entry, src, tmp_path)
        assert res["ok"], res.get("error")
        out = Path(res["output_path"])
        assert out.is_file() and out.stat().st_size > 0
        assert res["verification"]["ok"], res["verification"]
        independent(cap, out, src, tmp_path)
    except BaseException as exc:
        record(
            entry,
            "failed",
            time.time() - t0,
            {"error": f"{type(exc).__name__}: {str(exc)[:200]}"},
        )
        raise
    record(entry, "passed", time.time() - t0, {"fidelity": cap.fidelity.value})
