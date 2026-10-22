"""Deterministic generation of real input fixtures for every supported source format.

Files are produced by independent tools (Pillow, FFmpeg, Pandoc, LibreOffice, python-pptx,
openpyxl, py7zr, stdlib) and cached under tests/.fixture-cache with a cross-process lock so
parallel pytest workers do not regenerate them. Document fixtures carry SENTINEL so content
preservation can be asserted after conversion.
"""

from __future__ import annotations

import gzip
import io
import json
import os
import shutil
import subprocess
import tarfile
import tempfile
import zipfile
from pathlib import Path

try:
    import fcntl

    def _lock_exclusive(f) -> None:
        fcntl.flock(f, fcntl.LOCK_EX)

except ImportError:  # Windows: lock the first byte; the OS releases it when the file is closed
    import msvcrt

    def _lock_exclusive(f) -> None:
        f.seek(0)
        msvcrt.locking(f.fileno(), msvcrt.LK_LOCK, 1)


SENTINEL = "EZCONVERTER-SENTINEL-4711"
CACHE = Path(__file__).resolve().parent.parent / ".fixture-cache"

EXT = {"jpeg": "jpg", "tiff": "tif", "tgz": "tar.gz", "mpeg": "mpg"}


def _run(argv: list[str], cwd: Path | None = None) -> None:
    r = subprocess.run(argv, cwd=cwd, capture_output=True, text=True, timeout=300)
    if r.returncode != 0:
        raise RuntimeError(f"{argv[0]} failed: {r.stderr[-500:]}")


def _lo(src: Path, filt: str, outdir: Path) -> Path:
    prof = Path(tempfile.mkdtemp(prefix="lo-fx-"))
    try:
        _run(
            [
                "soffice",
                f"-env:UserInstallation={prof.absolute().as_uri()}",
                "--headless",
                "--nologo",
                "--norestore",
                "--convert-to",
                filt,
                "--outdir",
                str(outdir),
                str(src),
            ]
        )
    finally:
        shutil.rmtree(prof, ignore_errors=True)
    ext = filt.split(":")[0]
    return outdir / f"{src.stem}.{ext}"


MD = f"""# Quarterly report

This paragraph carries the sentinel {SENTINEL} for content checks.

| Item | Value |
|------|-------|
| alpha | 1 |
| beta | 2 |

- first point
- second point
"""


ARCHIVE_FILES = {
    "hello.txt": f"hello {SENTINEL}\n".encode(),
    "dir/data.json": b'{"k": [1, 2, 3]}',
    "dir/sub/blob.bin": bytes(range(256)) * 40,
}


def _make(fmt: str, out: Path) -> None:  # noqa: C901
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = Path(tempfile.mkdtemp(prefix="fx-"))
    try:
        if fmt in ("png", "jpeg", "webp", "avif", "gif", "tiff", "bmp", "ico"):
            from PIL import Image, ImageDraw

            im = Image.new("RGBA", (96, 64), (220, 30, 30, 255))
            ImageDraw.Draw(im).rectangle([48, 0, 95, 63], fill=(30, 30, 220, 255))
            if fmt in ("jpeg", "bmp"):
                im = im.convert("RGB")
            im.save(
                out,
                {"jpeg": "JPEG", "tiff": "TIFF", "ico": "ICO"}.get(fmt, fmt.upper()),
            )
        elif fmt == "svg":
            out.write_text(
                '<svg xmlns="http://www.w3.org/2000/svg" width="96" height="64" viewBox="0 0 96 64">'
                '<rect width="48" height="64" fill="#dc1e1e"/><rect x="48" width="48" height="64" fill="#1e1edc"/></svg>'
            )
        elif fmt in ("mp3", "wav", "flac", "aac", "m4a", "ogg", "opus", "aiff"):
            codec = {
                "opus": ["-c:a", "libopus"],
                "ogg": ["-c:a", "libvorbis"],
                "m4a": ["-c:a", "aac"],
            }.get(fmt, [])
            fmt_flag = ["-f", "adts"] if fmt == "aac" else []
            _run(
                [
                    "ffmpeg",
                    "-loglevel",
                    "error",
                    "-y",
                    "-f",
                    "lavfi",
                    "-i",
                    "sine=frequency=440:duration=2",
                    *codec,
                    *fmt_flag,
                    str(out),
                ]
            )
        elif fmt in ("mp4", "mkv", "mov", "avi", "webm", "mpeg"):
            vc = {
                "webm": ["-c:v", "libvpx", "-c:a", "libvorbis"],
                "avi": ["-c:v", "mpeg4", "-c:a", "libmp3lame"],
                "mpeg": ["-c:v", "mpeg2video", "-c:a", "mp2"],
                "mp4": ["-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac"],
                "mkv": ["-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac"],
                "mov": ["-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac"],
            }[fmt]
            _run(
                [
                    "ffmpeg",
                    "-loglevel",
                    "error",
                    "-y",
                    "-f",
                    "lavfi",
                    "-i",
                    "testsrc=duration=2:size=320x240:rate=15",
                    "-f",
                    "lavfi",
                    "-i",
                    "sine=frequency=440:duration=2",
                    "-shortest",
                    *vc,
                    str(out),
                ]
            )
        elif fmt in ("docx", "odt", "rtf", "html", "txt", "md", "pdf", "doc"):
            if fmt == "md":
                out.write_text(MD)
            elif fmt == "txt":
                out.write_text(
                    f"Quarterly report\n\nThis paragraph carries the sentinel {SENTINEL} for content checks.\n"
                )
            elif fmt == "html":
                out.write_text(
                    f"<!doctype html><html><head><title>t</title></head><body><h1>Quarterly report</h1><p>sentinel {SENTINEL}</p></body></html>"
                )
            elif fmt in ("docx", "odt", "rtf"):
                (tmp / "s.md").write_text(MD)
                _run(["pandoc", "-s", str(tmp / "s.md"), "-o", str(out)])
            else:  # pdf, doc via LibreOffice from docx
                docx = fixture("docx")
                res = _lo(docx, "pdf" if fmt == "pdf" else "doc:MS Word 97", tmp)
                shutil.copyfile(res, out)
        elif fmt in ("xlsx", "xls", "ods", "csv"):
            if fmt == "csv":
                out.write_text(f"name,qty,note\nwidget,3,{SENTINEL}\ngadget,5,ok\n")
            else:
                import openpyxl

                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Sheet1"
                ws.append(["name", "qty", "note"])
                ws.append(["widget", 3, SENTINEL])
                ws.append(["gadget", 5, "ok"])
                wb.save(tmp / "s.xlsx")
                if fmt == "xlsx":
                    shutil.copyfile(tmp / "s.xlsx", out)
                else:
                    res = _lo(
                        tmp / "s.xlsx",
                        "xls:MS Excel 97" if fmt == "xls" else "ods",
                        tmp,
                    )
                    shutil.copyfile(res, out)
        elif fmt in ("pptx", "ppt", "odp"):
            from pptx import Presentation

            prs = Presentation()
            s = prs.slides.add_slide(prs.slide_layouts[1])
            s.shapes.title.text = "Quarterly report"
            s.placeholders[1].text = f"Sentinel {SENTINEL}"
            prs.save(tmp / "s.pptx")
            if fmt == "pptx":
                shutil.copyfile(tmp / "s.pptx", out)
            else:
                res = _lo(
                    tmp / "s.pptx",
                    "ppt:MS PowerPoint 97" if fmt == "ppt" else "odp",
                    tmp,
                )
                shutil.copyfile(res, out)
        elif fmt in ("zip", "tar", "tgz", "7z"):
            files = ARCHIVE_FILES
            if fmt == "zip":
                with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
                    for n, b in files.items():
                        z.writestr(n, b)
            elif fmt in ("tar", "tgz"):
                with tarfile.open(out, "w:gz" if fmt == "tgz" else "w") as t:
                    for n, b in files.items():
                        ti = tarfile.TarInfo(n)
                        ti.size = len(b)
                        t.addfile(ti, io.BytesIO(b))
            else:
                import py7zr

                root = tmp / "r"
                for n, b in files.items():
                    (root / n).parent.mkdir(parents=True, exist_ok=True)
                    (root / n).write_bytes(b)
                with py7zr.SevenZipFile(out, "w") as z:
                    for n in files:
                        z.write(root / n, n)
        elif fmt == "gz":
            with gzip.open(out, "wb") as g:
                g.write(json.dumps({"sentinel": SENTINEL, "n": [1, 2, 3]}).encode())
        elif fmt == "json":
            out.write_text(
                json.dumps(
                    [
                        {"id": 1, "name": "widget", "tags": None, "note": SENTINEL},
                        {"id": 2, "name": "gadget", "tags": None, "note": "ok"},
                    ],
                    indent=1,
                )
            )
        elif fmt == "yaml":
            out.write_text(
                f"- id: 1\n  name: widget\n  note: {SENTINEL}\n- id: 2\n  name: gadget\n  note: ok\n"
            )
        elif fmt == "xml":
            out.write_text(
                f'<?xml version="1.0"?><catalog><item id="1"><name>widget</name><note>{SENTINEL}</note></item><item id="2"><name>gadget</name></item></catalog>'
            )
        else:
            raise KeyError(fmt)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def fixture(fmt: str) -> Path:
    """Return the path to a valid sample file of `fmt`, generating it once (cross-process safe)."""
    CACHE.mkdir(parents=True, exist_ok=True)
    out = CACHE / f"sample.{EXT.get(fmt, fmt)}"
    lock = CACHE / f".{fmt}.lock"
    with open(lock, "w") as lf:
        _lock_exclusive(lf)
        if not out.exists() or out.stat().st_size == 0:
            scratch = Path(tempfile.mkdtemp(prefix=".gen-", dir=CACHE))
            try:
                _make(
                    fmt, scratch / out.name
                )  # keep the real extension: ffmpeg/pandoc infer the format from it
                os.replace(scratch / out.name, out)
            finally:
                shutil.rmtree(scratch, ignore_errors=True)
    return out


MAGIC = {
    "png": b"\x89PNG\r\n\x1a\n",
    "jpeg": b"\xff\xd8\xff\xe0\x00\x10JFIF\x00",
    "gif": b"GIF89a",
    "webp": b"RIFF\x20\x00\x00\x00WEBPVP8 ",
    "bmp": b"BM" + b"\x00" * 12 + b"\x28\x00\x00\x00",
    "tiff": b"II*\x00\x08\x00\x00\x00",
    "pdf": b"%PDF-1.7\n",
    "mp3": b"ID3\x03\x00\x00\x00\x00\x00\x0a",
    "flac": b"fLaC",
    "wav": b"RIFF\x20\x00\x00\x00WAVEfmt ",
    "ogg": b"OggS\x00\x02",
    "mp4": b"\x00\x00\x00\x18ftypisom\x00\x00\x02\x00",
    "mkv": b"\x1a\x45\xdf\xa3",
    "avi": b"RIFF\x20\x00\x00\x00AVI LIST",
    "zip": b"PK\x03\x04\x14\x00",
    "gz": b"\x1f\x8b\x08\x00",
    "7z": b"7z\xbc\xaf\x27\x1c\x00\x04",
    "docx": b"PK\x03\x04\x14\x00",
    "doc": bytes.fromhex("D0CF11E0A1B11AE1"),
    "rtf": b"{\\rtf1\\ansi ",
}


def corrupt(fmt: str, dest: Path) -> Path:
    """A file that claims to be `fmt` (valid magic bytes) but whose body is garbage."""
    import random

    rng = random.Random(1234)
    dest.write_bytes(MAGIC[fmt] + bytes(rng.randrange(256) for _ in range(2048)))
    return dest
